import openpyxl
import pandas as pd
from pathlib import Path

excel_file = Path('vente.xlsx')
print(f"Reading: {excel_file}")
print(f"File exists: {excel_file.exists()}")

# Read with openpyxl to check structure
wb = openpyxl.load_workbook('vente.xlsx')
print(f"\nSheets: {wb.sheetnames}")

ws = wb.active
print(f"Active sheet: {ws.title}")
print(f"Max rows: {ws.max_row}")
print(f"Max columns: {ws.max_column}")

print("\n--- First 3 rows (header + sample) ---")
for i, row in enumerate(list(ws.iter_rows(values_only=True))[:3], 1):
    print(f"Row {i}: {row[:10]}")  # Print first 10 columns

# Read with pandas for analysis
print("\n--- DataFrame Info ---")
df = pd.read_excel('vente.xlsx', sheet_name=0)
print(f"Shape: {df.shape}")
print(f"\nColumns: {df.columns.tolist()}")
print(f"\nFirst row sample:")
print(df.iloc[0])
print(f"\nData types:")
print(df.dtypes)
print(f"\nBasic stats:")
print(df.describe())
